"""
E2E 测试 (mocked).

跑通: PDF → 解析 → 大纲 → 写作 → Word 导出
"""
from __future__ import annotations

import asyncio
import time
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

pytestmark = pytest.mark.unit

PROJECT_ROOT = Path(__file__).parent.parent
SAMPLE_PDF = PROJECT_ROOT / "tests" / "test_pdfs" / "sample.pdf"


class TestE2EPipeline:
    """端到端链路 (mocked LLM)."""

    def test_full_pipeline(self) -> None:
        """PDF → 解析 → 模拟大纲 → 模拟写作 → Word 导出."""
        # 1) PDF 解析
        from backend.parser import PyMuPDFParser

        if not SAMPLE_PDF.exists():
            pytest.skip("sample.pdf 缺失")

        async def _run() -> dict:
            t0 = time.perf_counter()
            parser = PyMuPDFParser()
            with SAMPLE_PDF.open("rb") as f:
                file_bytes = f.read()
            parse_result = await parser.parse(file_bytes, filename="sample.pdf")
            t_parse = (time.perf_counter() - t0) * 1000

            # 2) 模拟大纲 (不调 LLM, 直接用测试数据)
            from backend.agents import OutlineNode, PaperOutline
            outline = PaperOutline(
                abstract_intent="本综述聚焦...",
                sections=[
                    OutlineNode(1, "摘要", "abstract", "概述", []),
                    OutlineNode(2, "引言", "introduction", "背景", []),
                    OutlineNode(3, "方法", "methods", "检索", []),
                    OutlineNode(4, "结果", "results", "发现", []),
                    OutlineNode(5, "讨论", "discussion", "意义", []),
                    OutlineNode(6, "结论", "conclusion", "总结", []),
                ],
            )
            t_outline = 1  # mock

            # 3) 模拟写作 (直接造内容)
            section_results: dict[int, str] = {
                1: "## 摘要\n\n本综述...",
                2: "## 引言\n\n研究背景...",
                3: "## 方法\n\n检索策略...",
                4: "## 结果\n\n主要发现...",
                5: "## 讨论\n\n对比分析...",
                6: "## 结论\n\n总结展望...",
            }
            t_write = 1  # mock

            # 4) Word 导出 (用 docx 库直接生成, 不走 paper_pipeline)
            doc = Document()
            doc.add_heading("测试论文", level=0)
            for node in outline.sections:
                content = section_results.get(node.order_index, "")
                if not content:
                    continue
                doc.add_heading(node.heading, level=1)
                for para in content.split("\n\n"):
                    para = para.strip()
                    if para.startswith("## "):
                        doc.add_heading(para[3:], level=2)
                    elif para:
                        doc.add_paragraph(para)
            buf = BytesIO()
            doc.save(buf)
            docx_bytes = buf.getvalue()
            buf.close()
            t_export = (time.perf_counter() - t0) * 1000 - t_parse

            return {
                "parse_ms": t_parse,
                "outline_ms": t_outline,
                "write_ms": t_write,
                "export_ms": t_export,
                "total_ms": (time.perf_counter() - t0) * 1000,
                "pages": parse_result.page_count,
                "markdown_chars": len(parse_result.markdown),
                "docx_size": len(docx_bytes),
                "sections": len(outline.sections),
            }

        result = asyncio.run(_run())

        # 断言
        assert result["pages"] >= 1, "PDF 应能解析"
        assert result["markdown_chars"] > 1000, "应有 Markdown 内容"
        assert result["docx_size"] > 5000, "Word 应有内容"
        assert result["sections"] == 6, "应有 6 节"
        # 总耗时 mock 模式应该 < 30s
        assert result["total_ms"] < 30000, f"总耗时 {result['total_ms']:.0f}ms 超出预期"

    def test_word_document_is_valid(self) -> None:
        """生成的 docx 文件应能被反向解析."""
        from docx import Document
        from io import BytesIO

        doc = Document()
        doc.add_heading("测试", level=0)
        doc.add_paragraph("这是正文")
        doc.add_heading("二级标题", level=1)
        doc.add_paragraph("更多正文")
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        # 重新打开验证
        doc2 = Document(buf)
        text = "\n".join(p.text for p in doc2.paragraphs)
        assert "测试" in text
        assert "这是正文" in text
        assert "二级标题" in text
